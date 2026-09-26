"""PowerPoint generation like an agent's (python-pptx): title, bullets, tables, a picture and a
native chart per slide. usage: ppt_task.py <slides> <table rows> <seed> <out>"""
import io, random, sys
import matplotlib; matplotlib.use('Agg')
import matplotlib.pyplot as plt
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches, Pt
slides, trows, seed, out = int(sys.argv[1]), int(sys.argv[2]), int(sys.argv[3]), sys.argv[4]
random.seed(seed)
words = 'agent deploy build error python node server token cache memory report invoice release test merge'.split()
fig, ax = plt.subplots(figsize=(6, 3)); ax.hist([random.gauss(0, 1) for _ in range(2000)], bins=40)
img = io.BytesIO(); fig.savefig(img, format='png', dpi=100); plt.close(fig)
prs = Presentation()
s = prs.slides.add_slide(prs.slide_layouts[0]); s.shapes.title.text = 'Agent weekly review'; s.placeholders[1].text = 'generated'
for i in range(slides):
    s = prs.slides.add_slide(prs.slide_layouts[5]); s.shapes.title.text = f'Topic {i + 1}: ' + random.choice(words)
    tb = s.shapes.add_textbox(Inches(0.5), Inches(1.4), Inches(4.2), Inches(2)).text_frame
    for _ in range(5):
        para = tb.add_paragraph(); para.text = ' '.join(random.choice(words) for _ in range(8)); para.font.size = Pt(12)
    tbl = s.shapes.add_table(trows + 1, 4, Inches(0.5), Inches(3.6), Inches(4.2), Inches(2.5)).table
    for c, h in enumerate(['id', 'name', 'status', 'value']): tbl.cell(0, c).text = h
    for r in range(1, trows + 1):
        for c in range(4): tbl.cell(r, c).text = str(r) if c == 0 else random.choice(words) if c < 3 else f'{random.random():.3f}'
    img.seek(0); s.shapes.add_picture(img, Inches(5), Inches(1.4), width=Inches(4.5))
    cd = CategoryChartData(); cd.categories = ['Q1', 'Q2', 'Q3', 'Q4']
    cd.add_series('value', [random.randint(1, 100) for _ in range(4)])
    s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(5), Inches(4), Inches(4.5), Inches(2.5), cd)
prs.save(out)
print('ppt ok')
